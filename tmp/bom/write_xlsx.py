#!/usr/bin/env python3
"""Workbook rendering for the base-station bill of materials.

Kept separate from the data gathering so the sheet layout can change without
touching the schematic parsing or the distributor clients.
"""

from __future__ import annotations

import datetime
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TITLE_FONT = Font(bold=True, size=14)
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="44546A")
MANUAL_FILL = PatternFill("solid", fgColor="FFF2CC")  # still to be filled by hand
API_FILL = PatternFill("solid", fgColor="E2EFDA")  # came from a distributor API
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EURO = '#,##0.00\\ "€"'

COLUMNS = [
    ("Pos", 5),
    ("Menge", 7),
    ("Referenzen", 26),
    ("Wert (Schaltplan)", 20),
    ("Technik", 8),
    ("Footprint", 38),
    ("Hersteller", 18),
    ("Hersteller-Teilenr.", 24),
    ("Mouser-Teilenr.", 22),
    ("Mouser EUR", 11),
    ("Mouser Lager", 9),
    ("DigiKey-Teilenr.", 28),
    ("DigiKey EUR", 11),
    ("DigiKey Lager", 9),
    ("Lebenszyklus", 18),
    ("Bezugsquelle", 12),
    ("Einzelpreis EUR", 12),
    ("Gesamtpreis EUR", 13),
    ("Mouser", 9),
    ("DigiKey", 9),
    ("Datenblatt", 10),
    ("Status", 14),
    ("Hinweis", 60),
]
# Column numbers the formulas and fills refer to.
COL_QUANTITY, COL_MOUSER_PRICE, COL_DIGIKEY_PRICE = 2, 10, 13
COL_SOURCE, COL_UNIT, COL_TOTAL, COL_STATUS, COL_NOTE = 16, 17, 18, 22, 23
DISTRIBUTOR_COLUMNS = {"Mouser": (9, 10, 11, 19), "DigiKey": (12, 13, 14, 20)}


def _header(worksheet, title: str, subtitle: str, note: str, columns: list) -> int:
    worksheet["A1"] = title
    worksheet["A1"].font = TITLE_FONT
    worksheet["A2"] = subtitle
    worksheet["A3"] = note
    worksheet["A3"].font = Font(italic=True, size=9)
    row = 5
    for index, (label, width) in enumerate(columns, start=1):
        cell = worksheet.cell(row=row, column=index, value=label)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = worksheet.cell(row=row + 1, column=1)
    return row + 1


def _link(worksheet, row: int, column: int, label: str, url: str) -> None:
    if not url:
        return
    cell = worksheet.cell(row=row, column=column, value=label)
    cell.hyperlink = url
    cell.font = Font(color="0563C1", underline="single")


def _lifecycle(group: dict) -> str:
    """Worst lifecycle state reported by either distributor."""
    offers = [o for o in (group.get("offers") or {}).values() if o]
    ending = [o["lifecycle"] for o in offers if o["end_of_life"]]
    if ending:
        return ending[0]
    return next((o["lifecycle"] for o in offers if o["lifecycle"]), "")


def _bom_sheet(workbook, groups, technology, status_of, stamp) -> None:
    worksheet = workbook.create_sheet("Stückliste")
    priced = sum(1 for g in groups if g["price"] is not None)
    row = _header(
        worksheet,
        "Stückliste SmartHome-Basisstation",
        f"Quelle: PCB/BasisStation/*.kicad_sch · Stand {stamp} · "
        f"{len(groups)} Positionen · {sum(g['quantity'] for g in groups)} Bauteile",
        "Grün = Preis aus Mouser- bzw. DigiKey-API (Tagespreis, unverbindlich, Staffel passend zur "
        "Menge). Fett = günstigerer lieferbarer Preis. Gelb = von Hand zu füllen. Die Bezugsquelle "
        "steht auf dem günstigsten Anbieter mit ausreichend Lager und ist per Auswahlliste änderbar; "
        "Einzel-, Gesamtpreis und Summen rechnen sich daraus.",
        COLUMNS,
    )

    source_choice = DataValidation(type="list", formula1='"DigiKey,Mouser"', allow_blank=True)
    worksheet.add_data_validation(source_choice)

    first_data_row = row
    for position, group in enumerate(groups, start=1):
        status, note = status_of(group)
        offers = group.get("offers") or {}
        values = [
            position,
            group["quantity"],
            ", ".join(group["references"]),
            group["value"],
            technology(group["footprint"]),
            group["footprint"],
            group["manufacturer"],
            group["manufacturer_part"],
        ]
        for index, value in enumerate(values, start=1):
            worksheet.cell(row=row, column=index, value=value)

        deliverable_prices = {
            name: o["price"]
            for name, o in offers.items()
            if o and o["price"] is not None and o["stock"] >= group["quantity"]
        }
        cheapest = (
            min(deliverable_prices, key=deliverable_prices.get)
            if len(deliverable_prices) == 2
            and len(set(deliverable_prices.values())) == 2
            else ""
        )
        for name, (part_col, price_col, stock_col, link_col) in DISTRIBUTOR_COLUMNS.items():
            offer = offers.get(name)
            if not offer:
                continue
            worksheet.cell(row=row, column=part_col, value=offer["part"])
            price = worksheet.cell(row=row, column=price_col, value=offer["price"])
            if offer["price"] is not None:
                price.fill = API_FILL
                price.font = Font(bold=name == cheapest)
            stock = worksheet.cell(row=row, column=stock_col, value=offer["stock"])
            if offer["stock"] < group["quantity"]:
                stock.fill = WARN_FILL
            _link(worksheet, row, link_col, name, offer["url"])

        worksheet.cell(row=row, column=15, value=_lifecycle(group))
        source = worksheet.cell(row=row, column=COL_SOURCE, value=group["source"] or None)
        source_choice.add(source)
        if not group["source"]:
            source.fill = MANUAL_FILL
        elif not group.get("available"):
            source.fill = WARN_FILL

        src = f"{get_column_letter(COL_SOURCE)}{row}"
        mouser = f"{get_column_letter(COL_MOUSER_PRICE)}{row}"
        digikey = f"{get_column_letter(COL_DIGIKEY_PRICE)}{row}"
        unit = f"{get_column_letter(COL_UNIT)}{row}"
        quantity = f"{get_column_letter(COL_QUANTITY)}{row}"
        worksheet.cell(row=row, column=COL_UNIT).value = (
            f'=IF({src}="DigiKey",IF({digikey}="","",{digikey}),'
            f'IF({src}="Mouser",IF({mouser}="","",{mouser}),""))'
        )
        worksheet.cell(row=row, column=COL_TOTAL).value = (
            f'=IF(OR({unit}="",{quantity}=""),"",{quantity}*{unit})'
        )
        for column in (COL_MOUSER_PRICE, COL_DIGIKEY_PRICE, COL_UNIT, COL_TOTAL):
            worksheet.cell(row=row, column=column).number_format = EURO
        if group["price"] is None:
            worksheet.cell(row=row, column=COL_UNIT).fill = MANUAL_FILL

        _link(worksheet, row, 21, "PDF", group["datasheet"])
        worksheet.cell(row=row, column=COL_STATUS, value=status)
        worksheet.cell(row=row, column=COL_NOTE, value=note)
        if status in ("prüfen", "unvollständig", "Ersatz nötig"):
            worksheet.cell(row=row, column=COL_STATUS).fill = WARN_FILL
        elif status == "Vorschlag":
            worksheet.cell(row=row, column=COL_STATUS).fill = MANUAL_FILL

        for column in range(1, len(COLUMNS) + 1):
            cell = worksheet.cell(row=row, column=column)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, COL_NOTE))
        row += 1

    last_data_row = row - 1
    total_col = get_column_letter(COL_TOTAL)
    source_col = get_column_letter(COL_SOURCE)
    totals = [
        ("Summe erfasster Positionen", f"=SUM({total_col}{first_data_row}:{total_col}{last_data_row})"),
    ]
    for name in ("DigiKey", "Mouser"):
        totals.append(
            (
                f"davon {name}",
                f'=SUMIF({source_col}{first_data_row}:{source_col}{last_data_row},"{name}",'
                f"{total_col}{first_data_row}:{total_col}{last_data_row})",
            )
        )
    for offset, (label, formula) in enumerate(totals, start=1):
        worksheet.cell(row=row + offset, column=COL_UNIT - 1, value=label).font = Font(bold=offset == 1)
        total = worksheet.cell(row=row + offset, column=COL_TOTAL, value=formula)
        total.font = Font(bold=offset == 1)
        total.number_format = EURO
        total.border = BORDER

    counts = {n: sum(1 for g in groups if g["source"] == n) for n in ("DigiKey", "Mouser")}
    caveat = (
        "Die Summe ist deshalb eine Teilsumme, kein Gesamtpreis der Baugruppe."
        if priced < len(groups)
        else "Unter Vorbehalt: Vorschläge (Status gelb) sind noch nicht freigegeben, "
        "Positionen mit Status \"Ersatz nötig\" sind nicht ab Lager lieferbar."
    )
    worksheet.cell(
        row=row + len(totals) + 2,
        column=1,
        value=(
            f"Von {len(groups)} Positionen tragen {priced} einen Preis aus der API "
            f"({counts['DigiKey']} über DigiKey, {counts['Mouser']} über Mouser). " + caveat
        ),
    ).font = Font(italic=True, size=9)
    worksheet.auto_filter.ref = f"A{first_data_row - 1}:{get_column_letter(len(COLUMNS))}{last_data_row}"


def _single_sheet(workbook, instances, technology, stamp) -> None:
    worksheet = workbook.create_sheet("Einzelpositionen")
    columns = [
        ("Referenz", 12),
        ("Wert", 20),
        ("Technik", 8),
        ("Blatt", 30),
        ("Symbol (lib_id)", 30),
        ("Footprint", 40),
        ("Hersteller", 18),
        ("Hersteller-Teilenr.", 24),
        ("Mouser-Teilenr.", 22),
    ]
    row = _header(
        worksheet,
        "Einzelpositionen",
        f"Alle {len(instances)} bestückten Referenzen einzeln · Stand {stamp}",
        "Dient dem Abgleich gegen den Schaltplan — jede Referenz erscheint genau einmal.",
        columns,
    )
    for item in instances:
        values = [
            item["reference"],
            item["value"],
            technology(item["footprint"]),
            item["sheet"],
            item["lib_id"],
            item["footprint"],
            item["manufacturer"],
            f"{item['manufacturer_part']} (Vorschlag)" if item["proposal"] else item["manufacturer_part"],
            item["mouser_part"],
        ]
        for index, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row, column=index, value=value)
            cell.border = BORDER
        row += 1
    worksheet.auto_filter.ref = f"A5:I{row - 1}"


def _plural(count: int, singular: str, plural: str) -> str:
    """German needs the singular for exactly one; '1 Positionen' reads as a bug."""
    return f"{count} {singular if count == 1 else plural}"


def _open_points_sheet(workbook, groups, technology, stamp) -> None:
    worksheet = workbook.create_sheet("Offene Punkte")
    columns = [("Nr", 5), ("Betrifft", 34), ("Punkt", 52), ("Warum es zählt", 62), ("Wer entscheidet", 16)]
    generic = [
        g for g in groups if g["price"] is None and not g["mouser_part"] and not g["proposal"]
    ]
    proposed = [g for g in groups if g["proposal"]]
    through_hole = sorted(
        (r for g in groups if technology(g["footprint"]) == "THT" for r in g["references"]),
        key=lambda r: (re.match(r"([A-Za-z_]+)", r).group(1), int(re.sub(r"\D", "", r) or 0)),
    )
    row = _header(
        worksheet,
        "Offene Punkte vor der Bestellung",
        f"Stand {stamp} · Der Schaltplan ist elektrisch fertig (ERC 0 Fehler).",
        "Diese Punkte betreffen ausschließlich die Beschaffung, nicht die Funktion.",
        columns,
    )
    mismatched = [g for g in groups if g.get("mismatch")]
    end_of_life = [
        g for g in groups if any(o and o["end_of_life"] for o in (g.get("offers") or {}).values())
    ]
    unavailable = [
        g
        for g in groups
        if (g["mouser_part"] or g["manufacturer_part"])
        and not g.get("available")
        and g not in end_of_life
    ]
    single_source = [
        g
        for g in groups
        if g.get("available")
        and sum(
            1
            for o in (g.get("offers") or {}).values()
            if o and o["price"] is not None and o["stock"] >= g["quantity"]
        )
        == 1
    ]

    points = []
    if mismatched:
        points.append(
            (
                _plural(len(mismatched), "veraltete Mouser-Nummer im Schaltplan", "veraltete Mouser-Nummern im Schaltplan"),
                "; ".join(
                    f"{', '.join(g['references'])}: {g['schematic_mouser_part']} → {g['mouser_part']}"
                    for g in mismatched
                ),
                "Die im Symbol gepflegte Mouser-Nummer führt zu keinem Treffer mehr; die gültige "
                "Nummer stammt aus der Suche über die Herstellernummer. Solange das Schaltplanfeld "
                "nicht korrigiert ist, liefert jeder neue Export wieder die alte Nummer.",
                "Agent",
            )
        )
    if end_of_life:
        points.append(
            (
                _plural(len(end_of_life), "abgekündigte Position", "abgekündigte Positionen"),
                "; ".join(
                    f"{', '.join(g['references'])} ({g['manufacturer_part']}): "
                    + ", ".join(
                        f"{n} {o['lifecycle'] or 'ohne Statusangabe'}, Lager {o['stock']}"
                        for n, o in g["offers"].items()
                        if o
                    )
                    for g in end_of_life
                ),
                "Kein Beschaffungs-, sondern ein Abkündigungsproblem: auch ein Distributorwechsel "
                "hilft nicht. Ersatztyp wählen; hängt ein eigener Footprint daran, betrifft der "
                "Wechsel das Layout (Task 0041).",
                "Bediener",
            )
        )
    if unavailable:
        points.append(
            (
                _plural(len(unavailable), "Position nicht ab Lager lieferbar", "Positionen nicht ab Lager lieferbar"),
                "; ".join(f"{', '.join(g['references'])} ({g['value']})" for g in unavailable),
                "Teil ist bekannt, aber weder Mouser noch DigiKey haben es in der nötigen Menge "
                "ab Lager. Lieferzeit klären oder Alternative wählen, bevor der Termin für den "
                "Prototypenaufbau steht.",
                "Bediener",
            )
        )
    if single_source:
        points.append(
            (
                _plural(len(single_source), "Position nur bei einem Anbieter", "Positionen nur bei einem Anbieter"),
                "; ".join(f"{', '.join(g['references'])} (nur {g['source']})" for g in single_source),
                "Kein Ausweichen möglich, falls der Anbieter ausverkauft ist. Beim Wechsel auf "
                "ein breiter verfügbares Teil den Footprint im Blick behalten.",
                "Bediener",
            )
        )

    if proposed:
        points.append(
            (
                _plural(len(proposed), "Position mit Teilevorschlag", "Positionen mit Teilevorschlag"),
                "; ".join(f"{', '.join(g['references'])}: {g['manufacturer_part']}" for g in proposed),
                "Der Schaltplan trägt hier nur einen Wert oder ein abgekündigtes Teil. Die Vorschläge "
                "folgen einem Hausstandard (Widerstände 0805 1 %, Kondensatoren X7R 50 V bzw. "
                "X5R/X7R 16–25 V ab 1 µF, C0G am Quarz) und sind in Spalte Status gelb markiert. "
                "Nach Freigabe als Felder in den Schaltplan übernehmen, damit der nächste Export "
                "sie ohne Skript-Tabelle trägt. J1/J_PWR1: Zeichnung des Nachfolgers vor der "
                "Übernahme gegen den Footprint prüfen.",
                "Bediener",
            )
        )
    if generic:
        points.append(
            (
                _plural(len(generic), "Position ohne Teilenummer", "Positionen ohne Teilenummer"),
                "; ".join(f"{', '.join(g['references'])} ({g['value']})" for g in generic),
                "Für eine Bestellung fehlen Toleranz, Spannungsfestigkeit, Dielektrikum bzw. "
                "Belastbarkeit.",
                "Bediener",
            )
        )

    points += [
        (
            "Abgeleitete Teilenummern",
            "Bei D11, D12, U1, L1 und J6 stammt die Nummer aus Wert oder Bibliotheksnamen, "
            "nicht aus einem gepflegten Feld.",
            "In der Spalte Status als \"abgeleitet\" markiert. Vor der Bestellung am Datenblatt "
            "bestätigen — Serienbezeichnungen wie SRN6045TA brauchen noch den Wert-Suffix.",
            "Bediener",
        ),
        (
            "F1 (Sicherung)",
            "Der Wert im Schaltplan lautet schlicht \"Fuse\" — kein Strom- oder Auslösewert.",
            "Nennstrom steht in power_supply.md (Polyfuse 1,1 A); der Vorschlag MF-NSMF110-2 "
            "folgt dem. Den Wert auch im Schaltplan eintragen, sonst liest man dort nur \"Fuse\".",
            "Agent",
        ),
        (
            "FB1 (Ferritperle)",
            "Nutzt das Symbol Device:R — also einen Widerstand.",
            "Elektrisch folgenlos (zweipoliges Passivteil), aber Symbol und Bauteil "
            "widersprechen sich. Wer den Schaltplan liest, sieht einen Widerstand.",
            "Agent",
        ),
        (
            "J6 (Stiftleiste)",
            "Wert ist der generische KiCad-Symbolname \"Conn_02x03_Odd_Even\".",
            "Der Footprint verweist auf Würth 61200621621. Der Wert sollte die "
            "Bestellbezeichnung tragen, sonst ist die Stückliste ohne Blick in den "
            "Footprint nicht lesbar.",
            "Agent",
        ),
        (
            "D8, D9, D10 (LEDs)",
            "Nutzen das generische Symbol Device:LED, während D7 das spezifische "
            "WL-TMRC_3MM-Symbol nutzt.",
            "Uneinheitlich. Zudem liegen D8/D9 auf der Serie WL-TMRW, D10 auf WL-TMRC — "
            "zwei verschiedene Würth-Serien in derselben Funktionsgruppe.",
            "Bediener",
        ),
        (
            f"Mischbestückung THT / SMD ({len(through_hole)} bedrahtete Teile)",
            "Bedrahtet: " + ", ".join(through_hole) + ".",
            "Bedrahtete Teile erzwingen einen zweiten Bestückungsdurchlauf und verteuern die "
            "Fertigung spürbar. R13–R16 und R23–R26 sitzen zudem im Footprint DIN0617 "
            "(17 mm, 2-W-Klasse) — für 220 Ω und 4,7 kΩ in einer Signalfunktion "
            "überdimensioniert. Bei den LEDs und Steckverbindern ist THT bewusst gewählt; "
            "zu prüfen ist vor allem, ob Q1–Q4 und die Widerstände SMD werden können.",
            "Bediener",
        ),
        (
            "Preise sind Tagespreise",
            "Die grün hinterlegten Einzelpreise stammen aus der Mouser- und der DigiKey-API "
            "zum Stand oben.",
            "Sie sind unverbindlich, gelten für die Staffel passend zur Menge einer Baugruppe "
            "und verstehen sich ohne Zoll, Versand und Steuer. Für ein Angebot neu abfragen.",
            "Bediener",
        ),
    ]
    for number, (affects, point, why, who) in enumerate(points, start=1):
        for index, value in enumerate([number, affects, point, why, who], start=1):
            cell = worksheet.cell(row=row, column=index, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=index in (2, 3, 4))
        row += 1


def write_workbook(path, groups, instances, technology, status_of) -> None:
    stamp = datetime.date.today().isoformat()
    workbook = Workbook()
    workbook.remove(workbook.active)
    _bom_sheet(workbook, groups, technology, status_of, stamp)
    _single_sheet(workbook, instances, technology, stamp)
    _open_points_sheet(workbook, groups, technology, stamp)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
